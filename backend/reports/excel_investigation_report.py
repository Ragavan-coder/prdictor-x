"""
PREDICTRA-X Phase 9: Multi-Sheet Excel Forensic Audit Workbook Generator
Uses openpyxl to generate comprehensive enterprise audit spreadsheets:
- Sheet 1: FDR Summary
- Sheet 2: Extracted Fields & Coordinates
- Sheet 3: Verification Results (4 Layers)
- Sheet 4: Anomaly Findings
- Sheet 5: Duplicate Records
- Sheet 6: Lifecycle Events
- Sheet 7: Audit History
"""
import io
from datetime import datetime, timezone
from typing import Dict, Any, List

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def generate_fdr_investigation_excel(
    fdr_data: Dict[str, Any],
    verification_data: Dict[str, Any],
    extracted_fields: List[Dict[str, Any]],
    anomalies: List[Dict[str, Any]],
    duplicates: List[Dict[str, Any]],
    lifecycle_events: List[Dict[str, Any]],
    audit_events: List[Dict[str, Any]],
) -> io.BytesIO:
    """
    Generates a forensic investigation workbook in XLSX format.
    Returns a BytesIO stream ready for download.
    """
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    header_fill = PatternFill(start_color="C8102E", end_color="C8102E", fill_type="solid")
    sub_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True, color="0F172A")
    regular_font = Font(name="Calibri", size=10, color="0F172A")
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    # --------------------------------------------------------------------------
    # 1. Sheet: FDR Summary
    # --------------------------------------------------------------------------
    ws1 = wb.create_sheet(title="FDR Summary")
    ws1.append(["PREDICTRA-X FIXED DEPOSIT FORENSIC AUDIT DOSSIER"])
    ws1.merge_cells("A1:D1")
    ws1["A1"].font = Font(name="Calibri", size=14, bold=True, color="C8102E")
    ws1["A1"].alignment = Alignment(horizontal="center")
    ws1.append([])

    summary_rows = [
        ("FDR Reference Number", fdr_data.get("receipt_number", "N/A")),
        ("Issuing Bank", fdr_data.get("bank_name", "N/A")),
        ("Branch Name", fdr_data.get("branch_name", "N/A")),
        ("IFSC Code", fdr_data.get("ifsc_code", "N/A")),
        ("Account Holder", fdr_data.get("account_holder", "N/A")),
        ("Principal Deposit Amount", f"₹{float(fdr_data.get('principal_amount', 0)):,.2f}"),
        ("Contractual Interest Rate", f"{fdr_data.get('interest_rate', 0)}%"),
        ("Tenure Duration", f"{fdr_data.get('tenure_months', 12)} Months"),
        ("Compounding Frequency", fdr_data.get("compounding_frequency", "Quarterly")),
        ("Issue Date", fdr_data.get("issue_date", "N/A")),
        ("Maturity Date", fdr_data.get("maturity_date", "N/A")),
        ("Maturity Amount", f"₹{float(fdr_data.get('maturity_amount', 0)):,.2f}"),
        ("Overall Verification Status", verification_data.get("overall_status", "REVIEW_REQUIRED")),
        ("Composite Risk Score", f"{verification_data.get('risk_score', 25.0)}/100 ({verification_data.get('risk_severity', 'LOW')})"),
        ("Bank Confirmation Status", verification_data.get("bank_confirmation_status", "PENDING")),
        ("Generated At", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")),
    ]

    for k, v in summary_rows:
        ws1.append([k, v])
        r = ws1.max_row
        ws1.cell(row=r, column=1).font = bold_font
        ws1.cell(row=r, column=1).fill = alt_fill
        ws1.cell(row=r, column=2).font = regular_font
        ws1.cell(row=r, column=1).border = thin_border
        ws1.cell(row=r, column=2).border = thin_border

    # --------------------------------------------------------------------------
    # 2. Sheet: Extracted Fields
    # --------------------------------------------------------------------------
    ws2 = wb.create_sheet(title="Extracted Fields")
    headers2 = ["Field Key", "Field Label", "Original Value", "Current / Reviewed Value", "Confidence", "Bounding Box Coordinates"]
    ws2.append(headers2)
    for col_idx in range(1, len(headers2) + 1):
        cell = ws2.cell(row=1, column=col_idx)
        cell.font = white_font
        cell.fill = header_fill

    for f in extracted_fields:
        ws2.append([
            f.get("field_key", ""),
            f.get("field_label", ""),
            str(f.get("original_value", "")),
            str(f.get("current_value", "")),
            f"{float(f.get('ocr_confidence', 0.95))*100:.1f}%",
            str(f.get("bounding_box", "")),
        ])
        for col_idx in range(1, len(headers2) + 1):
            ws2.cell(row=ws2.max_row, column=col_idx).border = thin_border
            ws2.cell(row=ws2.max_row, column=col_idx).font = regular_font

    # --------------------------------------------------------------------------
    # 3. Sheet: Verification Results
    # --------------------------------------------------------------------------
    ws3 = wb.create_sheet(title="Verification Results")
    headers3 = ["Layer #", "Layer Name", "Status", "Finding Type", "Severity", "Explanation", "Detection Method"]
    ws3.append(headers3)
    for col_idx in range(1, len(headers3) + 1):
        cell = ws3.cell(row=1, column=col_idx)
        cell.font = white_font
        cell.fill = sub_fill

    findings = verification_data.get("findings", [])
    for fd in findings:
        ws3.append([
            fd.get("layer_number", 1),
            f"Layer {fd.get('layer_number', 1)}",
            fd.get("severity", "LOW"),
            fd.get("finding_type", ""),
            fd.get("severity", "LOW"),
            fd.get("explanation", ""),
            fd.get("detection_method", ""),
        ])
        for col_idx in range(1, len(headers3) + 1):
            ws3.cell(row=ws3.max_row, column=col_idx).border = thin_border
            ws3.cell(row=ws3.max_row, column=col_idx).font = regular_font

    # --------------------------------------------------------------------------
    # 4. Sheet: Anomaly Findings
    # --------------------------------------------------------------------------
    ws4 = wb.create_sheet(title="Anomaly Findings")
    headers4 = ["Anomaly Type", "Target Field", "Severity", "Original Value", "Detected Pattern", "Reason for Flagging", "Recommended Action"]
    ws4.append(headers4)
    for col_idx in range(1, len(headers4) + 1):
        cell = ws4.cell(row=1, column=col_idx)
        cell.font = white_font
        cell.fill = header_fill

    for anom in anomalies:
        ws4.append([
            anom.get("anomaly_type", ""),
            anom.get("field_name", ""),
            anom.get("severity", "MEDIUM"),
            str(anom.get("original_value", "")),
            str(anom.get("detected_value_or_pattern", "")),
            anom.get("reason_for_flagging", ""),
            anom.get("recommended_action", ""),
        ])
        for col_idx in range(1, len(headers4) + 1):
            ws4.cell(row=ws4.max_row, column=col_idx).border = thin_border
            ws4.cell(row=ws4.max_row, column=col_idx).font = regular_font

    # --------------------------------------------------------------------------
    # 5. Sheet: Duplicate Records
    # --------------------------------------------------------------------------
    ws5 = wb.create_sheet(title="Duplicate Records")
    headers5 = ["Match Tier", "Duplicate Reference", "Similarity Score", "Matched Fields", "Conflicting Fields", "Recommended Action"]
    ws5.append(headers5)
    for col_idx in range(1, len(headers5) + 1):
        cell = ws5.cell(row=1, column=col_idx)
        cell.font = white_font
        cell.fill = sub_fill

    for d in duplicates:
        ws5.append([
            d.get("match_level", ""),
            d.get("duplicate_fdr_ref") or d.get("primary_fdr_ref") or "N/A",
            f"{float(d.get('similarity_score', 1.0))*100:.1f}%",
            ", ".join(d.get("matched_fields", [])),
            ", ".join(d.get("conflicting_fields", [])),
            d.get("recommended_action", ""),
        ])
        for col_idx in range(1, len(headers5) + 1):
            ws5.cell(row=ws5.max_row, column=col_idx).border = thin_border
            ws5.cell(row=ws5.max_row, column=col_idx).font = regular_font

    # --------------------------------------------------------------------------
    # 6. Sheet: Lifecycle Events
    # --------------------------------------------------------------------------
    ws6 = wb.create_sheet(title="Lifecycle Events")
    headers6 = ["Event Type", "Event Description", "Performed Date", "Performed By", "Timestamp"]
    ws6.append(headers6)
    for col_idx in range(1, len(headers6) + 1):
        cell = ws6.cell(row=1, column=col_idx)
        cell.font = white_font
        cell.fill = header_fill

    for le in lifecycle_events:
        ws6.append([
            le.get("event_type", ""),
            le.get("event_description", ""),
            le.get("performed_date", ""),
            le.get("performed_by", ""),
            le.get("created_at", ""),
        ])
        for col_idx in range(1, len(headers6) + 1):
            ws6.cell(row=ws6.max_row, column=col_idx).border = thin_border
            ws6.cell(row=ws6.max_row, column=col_idx).font = regular_font

    # --------------------------------------------------------------------------
    # 7. Sheet: Audit History
    # --------------------------------------------------------------------------
    ws7 = wb.create_sheet(title="Audit History")
    headers7 = ["Block #", "Action", "Actor", "Resource ID", "Previous Hash", "Entry Hash", "Timestamp"]
    ws7.append(headers7)
    for col_idx in range(1, len(headers7) + 1):
        cell = ws7.cell(row=1, column=col_idx)
        cell.font = white_font
        cell.fill = sub_fill

    for ae in audit_events:
        ws7.append([
            ae.get("index", 0),
            ae.get("action", ""),
            ae.get("actor_username", ""),
            ae.get("resource_id", ""),
            ae.get("prev_hash", "")[:16] + "...",
            ae.get("entry_hash", "")[:16] + "...",
            ae.get("iso_time", ""),
        ])
        for col_idx in range(1, len(headers7) + 1):
            ws7.cell(row=ws7.max_row, column=col_idx).border = thin_border
            ws7.cell(row=ws7.max_row, column=col_idx).font = regular_font

    # Auto-adjust column widths on all sheets
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
